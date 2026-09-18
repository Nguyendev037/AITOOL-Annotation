#!/usr/bin/env python3
"""Sync guidelines from PDF/DOCX sources to Markdown using DeepSeek LLM.

Usage:
    python tools/sync_guidelines.py
    python tools/sync_guidelines.py --source guildlline --output guideline/semantic_segmentation.md
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from pathlib import Path

SUPPORTED_SUFFIXES = {".docx", ".pdf"}


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader
    return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def _extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs]
    tables = []
    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        tables.append("\n".join(rows))
    return "\n".join(paragraphs + tables)


def extract_sources(source_dir: Path) -> str:
    files = sorted(p for p in source_dir.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES)
    if not files:
        raise RuntimeError(f"No PDF or DOCX files in: {source_dir}")
    sections = []
    for path in files:
        text = _extract_pdf(path) if path.suffix.lower() == ".pdf" else _extract_docx(path)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        if text:
            sections.append(f"SOURCE: {path.relative_to(source_dir).as_posix()}\n{text}")
    return "\n\n---\n\n".join(sections)


def rewrite_with_llm(source_text: str, *, api_key: str, base_url: str, model: str) -> str:
    from openai import OpenAI
    system = (
        "You convert annotation documents into one precise Markdown guideline for an AI "
        "model service. Preserve every rule, class name, distinction, and quality checklist. "
        "Write in Vietnamese where the source is Vietnamese. Return Markdown only."
    )
    client = OpenAI(api_key=api_key, base_url=base_url.rstrip("/"), timeout=120.0, max_retries=0)
    response = client.chat.completions.create(
        model=model, temperature=0, max_tokens=8192,
        messages=[{"role": "system", "content": system},
                  {"role": "user", "content": f"DOCUMENTS:\n\n{source_text}"}],
    )
    content = (response.choices[0].message.content or "").strip()
    if content.startswith("```") and content.endswith("```"):
        content = re.sub(r"^```(?:markdown)?\s*|\s*```$", "", content, flags=re.IGNORECASE).strip()
    return content + "\n"


def main():
    parser = argparse.ArgumentParser(description="Sync guidelines from PDF/DOCX to Markdown")
    parser.add_argument("--source", default="guildlline", help="Source directory with PDF/DOCX files")
    parser.add_argument("--output", default="guideline/semantic_segmentation.md", help="Output markdown file")
    args = parser.parse_args()

    root = Path(__file__).resolve().parents[1]
    source_dir = root / args.source
    output_file = root / args.output

    api_key = os.getenv("DEEPSEEK_API_KEY", "")
    base_url = os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com")
    model = os.getenv("DEEPSEEK_MODEL", "deepseek-chat")

    if not api_key:
        print("ERROR: DEEPSEEK_API_KEY not set", file=sys.stderr)
        sys.exit(1)

    print(f"Extracting from: {source_dir}")
    source_text = extract_sources(source_dir)
    print(f"Rewriting with {model}...")
    markdown = rewrite_with_llm(source_text, api_key=api_key, base_url=base_url, model=model)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    output_file.write_text(markdown, encoding="utf-8")
    print(f"Written to: {output_file}")


if __name__ == "__main__":
    main()
